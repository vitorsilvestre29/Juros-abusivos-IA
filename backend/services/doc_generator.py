"""
Gerador de documentos jurídicos profissionais.

Utiliza scripts Node.js (docx-js) para produzir arquivos .docx com
formatação de escritório — fontes, margens ABNT, cabeçalho/rodapé
com número de página, tabelas de metadados, blocos de assinatura etc.

Fallback: python-docx básico caso o Node.js não esteja disponível.
"""

import json
import os
import subprocess
import io
from datetime import datetime
from typing import Optional

# ── Caminho dos scripts Node.js ───────────────────────────────────────────────
_SCRIPTS_DIR = os.path.join(os.path.dirname(__file__), '..', 'docx_scripts')
_NODE_BIN    = 'node'


def _run_node_script(script_name: str, payload: dict) -> Optional[bytes]:
    """
    Executa um script Node.js passando `payload` como JSON via stdin.
    Retorna os bytes do .docx gerado, ou None em caso de erro.
    """
    script_path = os.path.join(_SCRIPTS_DIR, script_name)
    if not os.path.isfile(script_path):
        return None
    try:
        result = subprocess.run(
            [_NODE_BIN, script_path],
            input=json.dumps(payload, ensure_ascii=False).encode('utf-8'),
            capture_output=True,
            timeout=60,
        )
        if result.returncode == 0 and result.stdout:
            return result.stdout
        # Log stderr se houver erro
        if result.stderr:
            print(f"[doc_generator] node stderr ({script_name}):", result.stderr.decode('utf-8', errors='replace')[:500])
        return None
    except (subprocess.TimeoutExpired, FileNotFoundError, Exception) as e:
        print(f"[doc_generator] Erro ao rodar {script_name}: {e}")
        return None


# ─────────────────────────────────────────────────────────────────────────────
# API PÚBLICA
# ─────────────────────────────────────────────────────────────────────────────

def generate_parecer_docx(case_data: dict, parecer_text: str) -> bytes:
    """Gera o Parecer Técnico Jurídico em .docx (Node.js → fallback python-docx)."""
    payload = {
        'client_name':    case_data.get('client_name', ''),
        'client_cpf':     case_data.get('client_cpf', ''),
        'client_address': case_data.get('client_address', ''),
        'case_type':      case_data.get('case_type', ''),
        'parecer_text':   parecer_text,
        'date_str':       datetime.now().strftime('%d/%m/%Y'),
    }
    result = _run_node_script('parecer.js', payload)
    if result:
        return result
    # ── Fallback python-docx ────────────────────────────────────────────────
    return _fallback_parecer(case_data, parecer_text)


def generate_procuracao_docx(case_data: dict, escritorio_data: dict = None) -> bytes:
    """Gera a Procuração Ad Judicia em .docx (Node.js → fallback python-docx)."""
    payload = {
        'client_name':        case_data.get('client_name', ''),
        'client_cpf':         case_data.get('client_cpf', ''),
        'client_rg':          case_data.get('client_rg', ''),
        'client_address':     case_data.get('client_address', ''),
        'client_nationality': case_data.get('client_nationality', 'brasileiro(a)'),
        'client_marital':     case_data.get('client_marital', ''),
        'client_profession':  case_data.get('client_profession', ''),
        'date_str':           datetime.now().strftime('%d de %B de %Y'),
    }
    result = _run_node_script('procuracao.js', payload)
    if result:
        return result
    return _fallback_procuracao(case_data)


def generate_peticao_docx(case_data: dict, peticao_text: str) -> bytes:
    """Gera a Petição Inicial em .docx (Node.js → fallback python-docx)."""
    payload = {
        'client_name':    case_data.get('client_name', ''),
        'client_cpf':     case_data.get('client_cpf', ''),
        'client_address': case_data.get('client_address', ''),
        'case_type':      case_data.get('case_type', ''),
        'peticao_text':   peticao_text,
        'date_str':       datetime.now().strftime('%d/%m/%Y'),
    }
    result = _run_node_script('peticao.js', payload)
    if result:
        return result
    return _fallback_peticao(case_data, peticao_text)


def generate_relatorio_preliminar_pdf(case_data: dict, contracts_analysis: list[dict]) -> bytes:
    """Gera relatório preliminar em PDF com achados técnicos e aviso legal."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    y = height - 2 * cm

    def draw_line(text: str, step: float = 0.6 * cm, bold: bool = False):
        nonlocal y
        if y < 2 * cm:
            c.showPage()
            y = height - 2 * cm
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 10)
        c.drawString(2 * cm, y, text[:115])
        y -= step

    c.setTitle("Relatorio Preliminar")
    draw_line("RELATORIO PRELIMINAR DE ANALISE TECNICA", bold=True)
    draw_line(f"Data: {datetime.now().strftime('%d/%m/%Y')}")
    draw_line(f"Cliente: {case_data.get('client_name', 'Nao informado')}")
    draw_line(f"CPF: {case_data.get('client_cpf', 'Nao informado')}")
    draw_line(f"Tipo do caso: {case_data.get('case_type', 'Nao informado')}")
    draw_line("", step=0.3 * cm)

    total_estimated = 0.0
    for idx, analysis in enumerate(contracts_analysis, start=1):
        draw_line(f"Contrato {idx}: {analysis.get('banco_identificado', 'Banco nao identificado')}", bold=True)
        draw_line(f"Numero: {analysis.get('numero_contrato', 'N/I')}")
        draw_line(f"Taxa mensal: {analysis.get('taxa_mensal', 'N/I')} | CET anual: {analysis.get('cet_anual', 'N/I')}")

        irregularidades = analysis.get("irregularidades", []) or []
        draw_line(f"Irregularidades encontradas: {len(irregularidades)}")
        for irr in irregularidades[:8]:
            draw_line(f"- {irr.get('tipo', 'Irregularidade')} ({irr.get('fundamento_legal', 'base legal nao informada')})")

        impacto = analysis.get("impacto_financeiro", {}) or {}
        estimated = impacto.get("estimated_overcharge_brl")
        if isinstance(estimated, (int, float)):
            total_estimated += float(estimated)
            draw_line(f"Impacto estimado: R$ {estimated:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."))
        draw_line("", step=0.3 * cm)

    draw_line("CONCLUSAO PRELIMINAR", bold=True)
    draw_line(
        f"Impacto financeiro estimado total: R$ {total_estimated:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    )
    draw_line("Ha viabilidade tecnica para analise juridica especializada, conforme achados acima.")
    draw_line("", step=0.3 * cm)
    draw_line("AVISO LEGAL", bold=True)
    draw_line(
        "Este documento e um laudo tecnico. A interpretacao juridica final e eventual acao revisional "
    )
    draw_line("devem ser realizadas por advogado regularmente inscrito na OAB.")
    draw_line("")
    draw_line("LGPD: tratamento de dados para execucao do servico e exercicio regular de direitos.")

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer.getvalue()


# ─────────────────────────────────────────────────────────────────────────────
# FALLBACKS python-docx (caso Node.js não esteja disponível no servidor)
# ─────────────────────────────────────────────────────────────────────────────

def _fb_imports():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    import re
    return Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, qn, re


def _fallback_parecer(case_data: dict, parecer_text: str) -> bytes:
    Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, qn, re = _fb_imports()
    BLUE = RGBColor(0, 51, 102)
    GRAY = RGBColor(89, 89, 89)
    ROMANO = re.compile(r'^(I{1,3}V?|IV|VI{0,3}|VII|VIII|IX|X+)(\.[0-9]+)?\s*[—\-–\.]\s*\S', re.IGNORECASE)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1.18); sec.bottom_margin = Inches(0.79)
    sec.left_margin = Inches(1.18); sec.right_margin = Inches(0.79)

    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('PARECER TÉCNICO JURÍDICO'); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = BLUE

    doc.add_paragraph(f"Cliente: {case_data.get('client_name','')} | CPF: {case_data.get('client_cpf','')} | Data: {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph('─' * 80)

    for line in parecer_text.split('\n'):
        s = line.strip()
        if not s: doc.add_paragraph(); continue
        p = doc.add_paragraph()
        letters = ''.join(c for c in s if c.isalpha())
        all_caps = letters and letters == letters.upper() and len(s) < 100
        if ROMANO.match(s) or (all_caps and len(s) < 80):
            r = p.add_run(s); r.bold = True; r.font.color.rgb = BLUE; r.font.size = Pt(12)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.5)
            r = p.add_run(s); r.font.size = Pt(12)

    doc.add_paragraph(f"\n[Cidade/UF], {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph('\n_______________________________________________\n[Nome do Advogado]\nOAB/[Estado] nº [Número]')
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf.getvalue()


def _fallback_procuracao(case_data: dict) -> bytes:
    Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, qn, re = _fb_imports()
    BLUE = RGBColor(0, 51, 102)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1.18); sec.bottom_margin = Inches(0.79)
    sec.left_margin = Inches(1.18); sec.right_margin = Inches(0.79)

    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('PROCURAÇÃO AD JUDICIA ET EXTRA'); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = BLUE

    nome  = case_data.get('client_name', '__________________')
    cpf   = case_data.get('client_cpf', '___.___.___-__')
    rg    = case_data.get('client_rg', '__________________')
    nac   = case_data.get('client_nationality', 'brasileiro(a)')
    ecivil = case_data.get('client_marital', '__________________')
    prof  = case_data.get('client_profession', '__________________')
    end   = case_data.get('client_address', '__________________')

    texto = (
        f"OUTORGANTE: {nome}, {nac}, {ecivil}, {prof}, portador(a) do RG nº {rg} e inscrito(a) "
        f"no CPF sob o nº {cpf}, residente e domiciliado(a) em {end}.\n\n"
        f"OUTORGADO(S): [NOME DO ADVOGADO], OAB/[ESTADO] nº [NÚMERO], "
        f"com escritório em [ENDEREÇO DO ESCRITÓRIO].\n\n"
        f"PODERES: Confere amplos poderes AD JUDICIA ET EXTRA para propor ação revisional "
        f"de contrato bancário, repetição de indébito, requerer tutelas de urgência e "
        f"praticar todos os atos necessários ao mandato.\n\n"
        f"[Cidade/UF], {datetime.now().strftime('%d de %B de %Y')}.\n\n"
        f"_______________________________________________\n{nome}\nCPF: {cpf}\nOUTORGANTE\n\n"
        f"_______________________________________________\n[NOME DO ADVOGADO]\nOAB/[ESTADO] nº [NÚMERO]\nOUTORGADO"
    )
    for bloco in texto.split('\n\n'):
        p = doc.add_paragraph(bloco.strip())
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf.getvalue()


def _fallback_peticao(case_data: dict, peticao_text: str) -> bytes:
    Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, qn, re = _fb_imports()
    BLUE = RGBColor(0, 51, 102)
    ROMANO = re.compile(r'^(I{1,3}V?|IV|VI{0,3}|VII|VIII|IX|X+)(\.[0-9]+)?\s*[—\-–\.]\s*\S', re.IGNORECASE)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(1.18); sec.bottom_margin = Inches(0.79)
    sec.left_margin = Inches(1.18); sec.right_margin = Inches(0.79)

    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run('PETIÇÃO INICIAL'); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = BLUE
    doc.add_paragraph(f"Requerente: {case_data.get('client_name','')} | CPF: {case_data.get('client_cpf','')} | Data: {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph('─' * 80)

    for line in peticao_text.split('\n'):
        s = line.strip()
        if not s: doc.add_paragraph(); continue
        p = doc.add_paragraph()
        letters = ''.join(c for c in s if c.isalpha())
        all_caps = letters and letters == letters.upper() and len(s) < 120
        if ROMANO.match(s) or (all_caps and len(s) < 90):
            r = p.add_run(s); r.bold = True; r.font.color.rgb = BLUE; r.font.size = Pt(12)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.5)
            r = p.add_run(s); r.font.size = Pt(12)

    doc.add_paragraph(f"\n[Cidade/UF], {datetime.now().strftime('%d/%m/%Y')}")
    doc.add_paragraph('\n_______________________________________________\n[Nome do Advogado]\nOAB/[Estado] nº [Número]')
    buf = io.BytesIO(); doc.save(buf); buf.seek(0); return buf.getvalue()
